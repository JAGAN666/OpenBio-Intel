"""Tests for POST /api/export/excel workbook structure (via export_excel)."""

import io
import re
from datetime import datetime, timezone
from unittest.mock import patch

from openpyxl import load_workbook

from api import ExcelExportRequest, export_excel
from research_agent import TrialRow


def _minimal_row() -> TrialRow:
    return TrialRow(
        nct_id="NCT01234567",
        sponsor="Example Pharma",
        phase="Phase 3",
        interventions=["drug-a"],
        mechanism_or_findings="First-line NSCLC: drug-a vs placebo — Phase 3.",
        mechanism_described=True,
    )


def _load_workbook_from_response(response):
    return load_workbook(io.BytesIO(response.body))


def _summary_values(ws):
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 5)}


@patch("api.datetime")
def test_export_excel_summary_and_data_sheets_with_query(mock_datetime):
    fixed = datetime(2026, 3, 15, 12, 30, 0, tzinfo=timezone.utc)
    mock_datetime.now.return_value = fixed

    req = ExcelExportRequest(
        query="Which Phase 3 lung trials use pembrolizumab?",
        narrative_summary="One trial cited [NCT01234567].",
        table_data=[_minimal_row()],
    )
    resp = export_excel(req)
    wb = _load_workbook_from_response(resp)

    assert wb.sheetnames == ["Summary", "Clinical Trials"]
    summary = _summary_values(wb["Summary"])
    assert summary["Original Question"] == req.query
    assert summary["Narrative Summary"] == req.narrative_summary
    assert summary["Row Count"] == 1
    assert summary["Generated At"] == "2026-03-15 12:30:00 UTC"

    data_ws = wb["Clinical Trials"]
    assert data_ws.freeze_panes == "A2"
    assert data_ws.max_row == 2
    assert data_ws.cell(row=2, column=1).value == "NCT01234567"


def test_export_excel_summary_question_placeholder_without_query():
    req = ExcelExportRequest(
        narrative_summary="Brief answer.",
        table_data=[],
    )
    resp = export_excel(req)
    wb = _load_workbook_from_response(resp)
    summary = _summary_values(wb["Summary"])
    assert summary["Original Question"] == "(not provided)"
    assert summary["Row Count"] == 0
    generated = summary["Generated At"]
    assert isinstance(generated, str)
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2} UTC", generated)
